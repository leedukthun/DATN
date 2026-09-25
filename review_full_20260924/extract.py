from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
import json, zipfile, re

src=Path(r'E:\Downloads\CNTT_2022606983_LeDucThuan_BaoCao.docx')
out=Path(__file__).parent
doc=Document(src); blocks=[]; lines=[]
for i,el in enumerate(doc.element.body.iterchildren()):
    if el.tag.endswith('}p'):
        p=Paragraph(el,doc); txt=p.text
        math=[''.join(x.text or '' for x in m.iter() if x.tag.endswith('}t')) for m in el.xpath('.//m:oMath')]
        imgs=el.xpath('.//a:blip/@r:embed')
        b={'id':i,'kind':'paragraph','style':p.style.name,'text':txt,'images':imgs,'math':math}
        if txt.strip() or imgs or math:
            lines.append(f'[{i:04d}] ({p.style.name}) {txt}'+(f' [IMAGES {imgs}]' if imgs else '')+(f' [MATH {math}]' if math else ''))
    elif el.tag.endswith('}tbl'):
        t=Table(el,doc); b={'id':i,'kind':'table','rows':[[c.text for c in r.cells] for r in t.rows]}
        lines.append(f'[{i:04d}] TABLE\n'+'\n'.join(' | '.join(r) for r in b['rows']))
    else: continue
    blocks.append(b)
(out/'blocks.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2),encoding='utf-8')
(out/'report_all.txt').write_text('\n'.join(lines),encoding='utf-8')
with zipfile.ZipFile(src) as z:
    (out/'app_properties.xml').write_bytes(z.read('docProps/app.xml'))
    for n in z.namelist():
        if n.startswith('word/media/'):
            dest=out/'media'/Path(n).name;dest.parent.mkdir(exist_ok=True);dest.write_bytes(z.read(n))
rels={k:v.target_ref for k,v in doc.part.rels.items()}
(out/'rels.json').write_text(json.dumps(rels,ensure_ascii=False,indent=2),encoding='utf-8')
print('Blocks:',len(blocks),'Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables),'Sections:',len(doc.sections))
for b in blocks:
    if b['kind']=='paragraph' and (b['style'].startswith('Heading') or re.match(r'^(CHƯƠNG|Chương|TÀI LIỆU|KẾT LUẬN)',b['text'])):
        print(f'[{b["id"]:04d}] {b["style"]} {b["text"]}')
