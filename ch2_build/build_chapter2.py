import io
import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r'D:\DATN\he_thong_phat_hien_khong_doi_mu_VSCode_scrollfix')
SOURCE = Path(r'D:\DATN\CNTT_2022606983_LeDucThuan_BaoCao.docx')
OUTPUT = ROOT / 'Chuong_2_rut_gon_khoang_15_trang.docx'


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), fill)
    tc_pr.append(shade)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for edge, value in [('top', 70), ('bottom', 70), ('left', 90), ('right', 90)]:
        item = OxmlElement(f'w:{edge}')
        item.set(qn('w:w'), str(value))
        item.set(qn('w:type'), 'dxa')
        mar.append(item)
    tc_pr.append(mar)


def plain_math(s):
    replacements = {
        r'\theta': 'θ', r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ',
        r'\rho': 'ρ', r'\tau': 'τ', r'\Delta': 'Δ', r'\geq': '≥',
        r'\lor': '∨', r'\times': '×', r'\min': 'min', r'\max': 'max',
        r'\cdot': '·',
        r'\ldots': '…', r'\mathbf': '', r'\mathrm': '', r'\text': '',
    }
    for old, new in replacements.items():
        s = s.replace(old, new)
    return s.replace('{', '').replace('}', '').replace(r'\_', '_')


def add_inline(p, content):
    pattern = re.compile(r'(\\\(.+?\\\)|\*\*.+?\*\*|\*.+?\*|`.+?`)', re.DOTALL)
    pos = 0
    for match in pattern.finditer(content):
        if match.start() > pos:
            p.add_run(content[pos:match.start()])
        token = match.group(0)
        if token.startswith(r'\('):
            run = p.add_run(plain_math(token[2:-2]))
            run.font.name = 'Cambria Math'
        elif token.startswith('**'):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith('*'):
            p.add_run(token[1:-1]).italic = True
        else:
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(11)
        pos = match.end()
    if pos < len(content):
        p.add_run(content[pos:])


MATH_P = {
    '2.1': 203, '2.2': 204, '2.7': 304, '2.8': 316,
    '2.9': 324, '2.10': 360, '2.11': 375, '2.12': 380,
    '2.13': 385, '2.14': 394, '2.15': 401, '2.16': 425,
    '2.17': 428, '2.18': 435,
}
MATH_T = {'2.4': 4, '2.5': 7, '2.6': 8}


def original_math(number):
    if number in MATH_P:
        return source_doc.paragraphs[MATH_P[number]]._p.xpath('.//m:oMath')[0]
    if number in MATH_T:
        return source_doc.tables[MATH_T[number]]._tbl.xpath('.//m:oMath')[0]
    return None


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.59)
sec.page_height = Cm(27.94)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(3.5)
sec.right_margin = Cm(2.0)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(0)

for name, size, before, after in [
    ('Title', 16, 0, 8),
    ('Heading 1', 14, 7, 4),
    ('Heading 2', 14, 6, 3),
    ('Heading 3', 14, 5, 3),
]:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
title.paragraph_format.keep_with_next = True
title_run = title.add_run('CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ SỬ DỤNG')
title_run.bold = True
title_run.font.size = Pt(16)
intro = doc.add_paragraph(
    'Chương này trình bày cơ sở phát hiện đối tượng bằng YOLO, dữ liệu gán nhãn và phương pháp huấn luyện tinh chỉnh mô hình từ hai lớp lên ba lớp. Tiếp đó, chương mô tả cách xử lý ảnh, video, liên kết vùng đầu với xe máy, theo dõi đối tượng và xác nhận trạng thái không đội mũ. Các công nghệ triển khai và tiêu chí đánh giá được xác định để làm cơ sở thiết kế, thực nghiệm hệ thống.'
)
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
intro.paragraph_format.first_line_indent = Cm(1)

source_doc = Document(SOURCE)
blocks = []
for part in ('part_2122.json', 'part_2325.json'):
    payload = json.loads((ROOT / 'ch2_build' / part).read_text(encoding='utf-8'))
    blocks.extend(payload['blocks'])

prev_type = None
for block in blocks:
    typ = block['type']
    if typ == 'heading':
        level = block.get('level', 1)
        p = doc.add_paragraph(style={0: 'Title', 1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3'}.get(level, 'Heading 3'))
        add_inline(p, block['text'])
        if level == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif typ == 'paragraph':
        p = doc.add_paragraph()
        add_inline(p, block['text'])
        if block['text'].startswith('Bảng 2.3.'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.italic = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1)
            if prev_type == 'table':
                p.paragraph_format.space_before = Pt(4)
    elif typ == 'equation':
        formula = block['text']
        m = re.search(r'\\tag\{(2\.\d+)\}', formula)
        number = block.get('number') or (m.group(1) if m else None)
        if not number and formula.startswith('C='):
            math_node = source_doc.paragraphs[300]._p.xpath('.//m:oMath')[0]
        else:
            math_node = original_math(number)
        if number == '2.3':
            paragraph = source_doc.paragraphs[226]
            blip = paragraph._p.xpath('.//a:blip')[0]
            rid = blip.get(qn('r:embed'))
            part = paragraph.part.related_parts[rid]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.add_run().add_picture(io.BytesIO(part.blob), width=Cm(15.8))
        elif math_node is not None:
            t = doc.add_table(rows=1, cols=2)
            t.autofit = False
            t.columns[0].width = Cm(14.00)
            t.columns[1].width = Cm(2.09)
            for ci, cell in enumerate(t.rows[0].cells):
                cell.width = Cm(14.00 if ci == 0 else 2.09)
            left = t.cell(0, 0).paragraphs[0]
            left.alignment = WD_ALIGN_PARAGRAPH.CENTER
            left._p.append(deepcopy(math_node))
            right = t.cell(0, 1).paragraphs[0]
            right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if number:
                right.add_run(f'({number})')
            for cell in t.rows[0].cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(0)
        else:
            p = doc.add_paragraph(plain_math(re.sub(r'\\tag\{.*?\}', '', formula)))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
    elif typ == 'table':
        caption = block.get('caption')
        if caption:
            cp = doc.add_paragraph()
            add_inline(cp, caption)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_before = Pt(4)
            cp.paragraph_format.space_after = Pt(3)
            cp.paragraph_format.keep_with_next = True
            for run in cp.runs:
                run.font.size = Pt(12)
                run.font.italic = True
        rows = block['rows']
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = block.get('widths_cm')
        if not widths:
            widths = [16.09 / len(rows[0])] * len(rows[0])
        for i, val in enumerate(widths):
            table.columns[i].width = Cm(val)
        for j, text in enumerate(rows[0]):
            table.rows[0].cells[j].text = ''
            add_inline(table.rows[0].cells[j].paragraphs[0], str(text))
        for row in rows[1:]:
            cells = table.add_row().cells
            for j, text in enumerate(row):
                cells[j].text = ''
                add_inline(cells[j].paragraphs[0], str(text))
        for ri, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            no_split = OxmlElement('w:cantSplit')
            tr_pr.append(no_split)
            if ri == 0:
                tbl_header = OxmlElement('w:tblHeader')
                tbl_header.set(qn('w:val'), 'true')
                tr_pr.append(tbl_header)
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade_cell(cell, 'E8EDF3' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF'))
                set_cell_margins(cell)
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing = 1.1
                    para.paragraph_format.space_after = Pt(0)
                    for run in para.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(12)
                        run.font.bold = ri == 0
    elif typ == 'figure':
        idx = int(block['source_paragraph'])
        paragraph = source_doc.paragraphs[idx]
        blip = paragraph._p.xpath('.//a:blip')[0]
        rid = blip.get(qn('r:embed'))
        part = paragraph.part.related_parts[rid]
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(io.BytesIO(part.blob), width=Cm(15.8))
        cp = doc.add_paragraph(block['caption'])
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(6)
        for run in cp.runs:
            run.font.size = Pt(12)
            run.font.italic = True
    else:
        raise ValueError(f'Unknown block type: {typ}')
    prev_type = typ

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

doc.save(OUTPUT)
print(OUTPUT)
print('Blocks:', len(blocks))
