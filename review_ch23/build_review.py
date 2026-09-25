from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUTPUT=ROOT.parent/'Ra_soat_va_bo_cuc_lai_Chuong_2_3.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21)
sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8)
sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2.0)
sec.right_margin=Cm(2.0)
sec.footer_distance=Cm(.8)

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
    s=doc.styles[name]
    s.font.name='Times New Roman'
    s.font.color.rgb=RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().rFonts
    rf.set(qn('w:eastAsia'),'Times New Roman')
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    s.font.size=Pt(11.5)
    s.paragraph_format.space_after=Pt(4)
    s.paragraph_format.line_spacing=1.0
    s.paragraph_format.widow_control=True
for name,size,before,after in [('Title',19,0,10),('Heading 1',14,0,8),('Heading 2',12,7,3),('Heading 3',11.5,6,3)]:
    s=doc.styles[name]
    s.font.size=Pt(size)
    s.font.bold=True
    s.paragraph_format.space_before=Pt(before)
    s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True

for e in doc.styles.element.xpath('.//w:pBdr'):
    e.getparent().remove(e)

def runs(p,text):
    for part in re.split(r'(\*\*.*?\*\*)',text):
        if not part: continue
        r=p.add_run(part[2:-2] if part.startswith('**') else part)
        if part.startswith('**'): r.bold=True

def table(lines):
    rows=[[v.strip() for v in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',x) for x in r)]
    n=len(rows[0]); t=doc.add_table(rows=0,cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=[2.2,5.4,9.4] if n==3 else [4.0,4.2,4.4,4.4]
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for idx,data in enumerate(rows):
        row=t.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
        for cell,txt,w in zip(row.cells,data,widths):
            cell.width=Cm(w)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E7E6E6' if idx==0 else 'FFFFFF');tcpr.append(shd)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.03
            runs(p,txt)
            for r in p.runs:
                r.font.size=Pt(10.5 if n==4 else 11)
                if idx==0:r.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=(ROOT/'ra_soat_bo_cuc.md').read_text(encoding='utf-8').splitlines()
i=0
compact_ch2=False
while i<len(lines):
    l=lines[i].strip()
    if not l:i+=1;continue
    if l=='<!-- PAGE -->':
        doc.add_page_break();i+=1;continue
    if l.startswith('## '):
        compact_ch2=l.startswith('## 4 Bố cục mới')
    if l.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i]);i+=1
        table(rows);continue
    if l.startswith('# '):p=doc.add_paragraph(style='Title');runs(p,l[2:])
    elif l.startswith('## '):p=doc.add_paragraph(style='Heading 1');runs(p,l[3:])
    elif l.startswith('### '):p=doc.add_paragraph(style='Heading 2');runs(p,l[4:])
    elif l.startswith('- '):p=doc.add_paragraph(style='List Bullet');runs(p,l[2:])
    else:
        p=doc.add_paragraph();runs(p,l)
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if compact_ch2 and not l.startswith('## '):
        p.paragraph_format.space_after=Pt(1.5)
        if l.startswith('### '):p.paragraph_format.space_before=Pt(5)
    i+=1

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(10)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='Rà soát và bố cục lại chương 2 và chương 3'
doc.core_properties.subject='Cơ sở lý thuyết, thiết kế và triển khai hệ thống phát hiện không đội mũ bảo hiểm'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.save(OUTPUT)
print(OUTPUT)
print('Paragraphs',len(doc.paragraphs),'Tables',len(doc.tables))
