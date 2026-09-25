from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r'D:\DATN\he_thong_phat_hien_khong_doi_mu_VSCode_scrollfix\Chuong_1_rut_gon_8_9_trang.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.59)
section.page_height = Cm(27.94)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(3.5)
section.right_margin = Cm(2.0)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.space_after = Pt(4)

for name, size, before, after in [('Title', 16, 0, 10), ('Heading 1', 14, 10, 7), ('Heading 2', 14, 8, 5)]:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


def title(s):
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(s)


def h1(s):
    doc.add_paragraph(s, style='Heading 1')


def h2(s):
    doc.add_paragraph(s, style='Heading 2')


def para(s):
    p = doc.add_paragraph(s)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1)
    return p


def table(caption, headings, rows, widths):
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.font.italic = True
        r.font.size = Pt(12)
    t = doc.add_table(rows=1, cols=len(headings))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = 'Table Grid'
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, text in enumerate(headings):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E9EEF4' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF'))
            tc_pr.append(shade)
            mar = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement(f'w:{edge}')
                item.set(qn('w:w'), '80')
                item.set(qn('w:type'), 'dxa')
                mar.append(item)
            tc_pr.append(mar)
            for p2 in cell.paragraphs:
                p2.paragraph_format.line_spacing = 1.15
                p2.paragraph_format.space_after = Pt(0)
                for run in p2.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                    run.font.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


title('CHƯƠNG 1 TỔNG QUAN VỀ BÀI TOÁN PHÁT HIỆN NGƯỜI ĐI XE MÁY KHÔNG ĐỘI MŨ BẢO HIỂM')
para('Chương này xác định nhu cầu ghi nhận các trường hợp không đội mũ bảo hiểm từ ảnh, video giao thông; mô tả bài toán, phạm vi và yêu cầu của hệ thống. Trọng tâm là phân tích hạn chế của mô hình hai lớp helmet, no-helmet khi gặp người đi bộ, từ đó trình bày hướng bổ sung lớp xe máy, liên kết đối tượng và theo dõi trong video. Các giới hạn được nêu để làm căn cứ thiết kế và đánh giá ở những chương tiếp theo.')

h1('1.1. Bối cảnh và nhu cầu thực tiễn')
h2('1.1.1. Nhu cầu ghi nhận trường hợp không đội mũ bảo hiểm')
para('Xe máy là phương tiện đi lại phổ biến. Việc đội mũ bảo hiểm góp phần bảo vệ người tham gia giao thông và thể hiện ý thức chấp hành quy định. Để hỗ trợ tuyên truyền và giám sát, các trường hợp không đội mũ cần được ghi nhận theo địa điểm, thời gian và lưu lại minh chứng có thể kiểm tra. Trong một cảnh giao thông, người lái, người ngồi sau, người đi bộ và các phương tiện khác có thể cùng xuất hiện. Vì vậy, nhận biết một vùng đầu không có mũ chưa đủ để khẳng định đó là người điều khiển xe máy.')
para('Cách ghi nhận thủ công dựa vào quan sát trực tiếp hoặc xem lại ảnh, video. Người quan sát xác định xe máy, nhận biết người liên quan, kiểm tra trạng thái đội mũ, sau đó ghi thời điểm và tổng hợp số liệu. Cách này cho phép xem xét ngữ cảnh, chẳng hạn phân biệt người đi bộ đứng cạnh xe với người đang ngồi trên xe. Tuy nhiên, video dài hoặc nhiều đối tượng di chuyển đòi hỏi thời gian, sự tập trung và quy tắc đếm thống nhất. Một xe xuất hiện ở nhiều khung hình có thể bị tính lặp; việc lưu minh chứng từ nhiều tệp cũng làm tăng khối lượng công việc.')
para('Thị giác máy tính có thể tự động phát hiện đối tượng, đánh dấu những trường hợp cần xem lại và hỗ trợ tổng hợp kết quả. Hệ thống trong đề tài hướng đến xử lý dữ liệu người dùng tải lên, lưu ảnh minh chứng cùng thông tin nguồn, địa điểm và khung giờ. Nhờ đó, người sử dụng có thể kiểm tra kết quả, lựa chọn tư liệu tuyên truyền và nhận biết khu vực, thời điểm cần quan tâm trong phạm vi dữ liệu đã phân tích.')
h2('1.1.2. Ý nghĩa và giới hạn của số liệu thống kê')
para('Số trường hợp được ghi nhận không tự nó phản ánh mức độ chấp hành của toàn khu vực. Địa điểm được ghi hình lâu hơn hoặc có lưu lượng xe lớn hơn có thể có tổng số trường hợp cao hơn dù tỷ lệ không đội mũ không cao hơn. Vì vậy, khi so sánh cần nêu rõ thời lượng quan sát, số phương tiện được ghi nhận, điều kiện ghi hình và chất lượng phát hiện. Đề tài không thực hiện điều tra diện rộng; mọi nhận xét về khu vực và khung giờ chỉ áp dụng cho dữ liệu khảo sát, thực nghiệm của hệ thống.')

h1('1.2. Mô tả bài toán và phạm vi hệ thống')
h2('1.2.1. Đối tượng, đầu vào và đầu ra')
para('Bài toán đặt ra là phân tích ảnh hoặc video giao thông để hỗ trợ phát hiện người điều khiển xe máy không đội mũ bảo hiểm. Quy trình cần xác định các vùng đầu, trạng thái đội mũ và vị trí xe máy; tiếp đó xem xét quan hệ giữa vùng đầu và phương tiện trước khi ghi nhận. Với video, hệ thống còn phải theo dõi đối tượng qua các khung hình để xác nhận trạng thái và tránh tính nhiều lần cùng một trường hợp.')
para('Dữ liệu đầu vào gồm một hoặc nhiều ảnh, video do người dùng tải lên. Mỗi phiên phân tích được gắn với dự án, địa điểm, ngày và khung giờ. Ảnh cung cấp thông tin tại một thời điểm nên bước liên kết chủ yếu dựa vào vị trí không gian. Video cung cấp chuỗi khung hình, cho phép đối chiếu các lần xuất hiện và xác định thời điểm ghi nhận. Cả hai loại dữ liệu đều có thể chứa người đi bộ, người ngồi sau, nhiều xe đứng gần nhau, vùng đầu nhỏ hoặc phương tiện bị che khuất.')
table('Bảng 1.1. Các lớp đối tượng sử dụng trong đề tài', ['Lớp', 'Ý nghĩa', 'Vai trò xử lý'], [
    ('helmet', 'Vùng đầu có mũ bảo hiểm', 'Cung cấp trạng thái đội mũ'),
    ('no-helmet', 'Vùng đầu không có mũ bảo hiểm', 'Xác định vùng đầu cần kiểm tra quan hệ với xe'),
    ('bike', 'Xe máy', 'Cung cấp vị trí phương tiện để liên kết')
], [2.7, 5.5, 7.9])
para('Đầu ra gồm ảnh hoặc video có đánh dấu kết quả nhận diện; danh sách trường hợp được ghi nhận kèm ảnh minh chứng, độ tin cậy, tệp nguồn và thời điểm trong video; trạng thái xử lý; cùng số liệu tổng hợp theo địa điểm và thời gian. Mỗi bản ghi phải truy ngược được dữ liệu nguồn để người dùng kiểm tra. Với video, mã theo dõi hỗ trợ xác định những lần xuất hiện thuộc cùng một đối tượng trong phạm vi tệp đang xử lý.')
para('Đơn vị thống kê phải nhất quán. Số vùng đầu, số người và số xe là những đại lượng khác nhau vì một xe có thể chở nhiều người. Khi tính tỷ lệ, tử số và mẫu số cần dùng cùng đơn vị quan sát. Việc thống nhất quy tắc này từ lúc ghi nhận đến lúc kiểm thử giúp số liệu có thể đối chiếu với kết quả kiểm đếm thủ công.')
h2('1.2.2. Yêu cầu và phạm vi giải quyết')
para('Về nhận diện, hệ thống cần phát hiện ba lớp nêu trên và sử dụng thông tin xe máy để hạn chế ghi nhận nhầm người đi bộ. Khi chưa phát hiện được xe hoặc quan hệ giữa các vùng chưa rõ, kết quả cần được xem là thiếu căn cứ xác nhận. Về video, mã đối tượng và việc xác nhận trạng thái qua nhiều khung hình được sử dụng để giảm ảnh hưởng của dự đoán sai nhất thời và hạn chế đếm trùng. Điều kiện xác nhận cũng cần tránh làm bỏ sót các đối tượng chỉ xuất hiện trong thời gian ngắn.')
para('Ứng dụng web cần cho phép quản lý dự án, địa điểm, tệp và phiên phân tích; cập nhật tiến độ, hiển thị kết quả, lưu minh chứng và tổng hợp thống kê. Chất lượng hệ thống phải được kiểm chứng bằng dữ liệu gán nhãn hoặc đối chiếu thủ công, không chỉ dựa vào độ tin cậy mà mô hình trả về. Phạm vi đề tài là ảnh, video tải lên; chưa kết nối camera để giám sát trực tiếp. Hệ thống không nhận diện danh tính, biển số, chất lượng mũ hay việc cài quai, và không tự động đưa ra quyết định xử phạt.')
para('Mục tiêu của đề tài hướng đến người điều khiển xe máy, nhưng mô hình ba lớp không có nhãn phân biệt người lái và người ngồi sau. Vì vậy, một vùng đầu liên kết với xe mới chứng minh được quan hệ không gian phù hợp, chưa đủ để xác định vai trò người lái. Giới hạn này cần được thể hiện trong kết quả và đánh giá, đặc biệt khi xe chở nhiều người có trạng thái đội mũ khác nhau.')

h1('1.3. Các hướng tiếp cận và nghiên cứu liên quan')
h2('1.3.1. Từ quan sát thủ công đến phát hiện tự động')
para('Quan sát thủ công có ưu điểm là người xem có thể dùng vị trí, tư thế và các khung hình liền kề để nhận biết ngữ cảnh. Đây là phương pháp đối chiếu phù hợp khi xây dựng tập kiểm thử, nhưng khó mở rộng khi số lượng tệp lớn. Ngược lại, mô hình phát hiện đối tượng có thể xử lý dữ liệu nhanh và trả về hộp giới hạn, nhãn cùng độ tin cậy. Kết quả tự động vẫn cần quy tắc liên kết và kiểm tra vì một nhãn no-helmet chỉ mô tả trạng thái vùng đầu, không mô tả quan hệ của người đó với xe máy.')
para('Hướng phát hiện trực tiếp hai lớp helmet và no-helmet có đầu ra đơn giản, thuận tiện để xây dựng hệ thống thử nghiệm. Hạn chế xuất hiện khi mọi vùng no-helmet đều bị xem là trường hợp cần ghi nhận: người đi bộ không đội mũ cũng có thể được tính. Hướng phát hiện đồng thời vùng đầu và xe máy cho phép kiểm tra quan hệ giữa các đối tượng trước khi ghi nhận, đổi lại cần dữ liệu gán nhãn xe và bước xử lý bổ sung. Với video, theo dõi đối tượng còn cần thiết để liên kết các kết quả ở những khung hình khác nhau.')
h2('1.3.2. Một số nghiên cứu có liên quan')
para('Siebert và Lin nghiên cứu cách tự động ghi nhận việc sử dụng mũ bảo hiểm từ video, đồng thời xem xét xe máy, số lượng và vị trí người trên xe. Công trình gợi ý rằng trạng thái đội mũ cần được đặt trong ngữ cảnh phương tiện và người liên quan, thay vì chỉ đếm vùng đầu. Aboah và cộng sự sử dụng YOLOv8 cho bài toán phát hiện vi phạm đội mũ, chú trọng lựa chọn mẫu khi dữ liệu gán nhãn còn hạn chế. Kinh nghiệm này có ý nghĩa khi chuẩn bị tập dữ liệu mở rộng lớp xe máy: ảnh huấn luyện cần bao quát cảnh khó và tránh nhiều mẫu gần giống nhau.')
para('Choi và Greer kết hợp phương pháp phát hiện, phân loại đối tượng liên quan đến xe máy và trạng thái đội mũ; nghiên cứu cũng lưu ý ảnh hưởng của độ phân giải thấp, khả năng quan sát hạn chế. Dù phương pháp khác với đề tài, các công trình trên cùng cho thấy nhu cầu xử lý mối quan hệ giữa người và phương tiện, cũng như đánh giá theo tình huống thực tế. Đề tài chỉ tham khảo định hướng này, không trực tiếp áp dụng mô hình của các nghiên cứu đã nêu.')
h2('1.3.3. Lựa chọn hướng giải quyết')
para('Đề tài lựa chọn phát triển từ mô hình YOLO hai lớp có sẵn để tận dụng khả năng nhận diện trạng thái đội mũ. Lớp bike được bổ sung nhằm cung cấp vị trí xe máy; sau nhận diện, hệ thống liên kết vùng đầu với xe dựa trên quan hệ không gian. Ở video, tracking duy trì mã đối tượng và xác nhận trạng thái qua nhiều khung hình. Kết quả cuối cùng được lưu và trình bày trong ứng dụng web để phục vụ xem lại, thống kê. Hiệu quả cần đánh giá ở cả cấp mô hình và cấp hệ thống: mô hình có thể phát hiện đúng vùng đầu, nhưng hệ thống vẫn có thể báo nhầm nếu liên kết sai đối tượng.')

h1('1.4. Mô hình ban đầu và hướng phát triển của đề tài')
h2('1.4.1. Hạn chế của mô hình hai lớp')
para('Mô hình ban đầu nhận diện hai lớp helmet và no-helmet, trả về hộp giới hạn, nhãn và độ tin cậy của từng phát hiện. Qua thử nghiệm ban đầu, hệ thống có thể ghi nhận người đi bộ không đội mũ thành trường hợp không đội mũ khi đi xe máy. Nguyên nhân chính là đầu ra no-helmet được dùng trực tiếp để tạo bản ghi, trong khi mô hình chưa cung cấp vị trí xe máy. Một phát hiện có độ tin cậy cao vẫn có thể là vùng đầu của người đi bộ; chỉ tăng ngưỡng độ tin cậy không bổ sung được thông tin ngữ cảnh cần thiết.')
para('Cần phân biệt lỗi nhận diện đối tượng với lỗi suy luận ở cấp hệ thống. Vùng đầu người đi bộ được phát hiện là no-helmet có thể đúng theo nhãn mô hình, nhưng suy ra người đó điều khiển xe máy là không có căn cứ. Từ quan sát này, đề tài thay đổi điều kiện ghi nhận: trạng thái không đội mũ phải được xem xét cùng thông tin về phương tiện. Mô hình hai lớp được giữ làm phương án so sánh để xác định mức độ cải thiện và những đánh đổi của hướng phát triển.')
h2('1.4.2. Mở rộng mô hình thành ba lớp')
para('Mô hình phát triển gồm helmet, no-helmet và bike. Việc bổ sung bike là mở rộng nhóm đối tượng nhận diện, không phải đề xuất kiến trúc mạng mới. Dữ liệu huấn luyện cần có nhãn xe máy nhất quán với hai lớp vùng đầu, đồng thời chứa các tình huống có người đi bộ gần xe, nhiều xe trong một ảnh và đối tượng bị che khuất. Nếu chỉ tập trung vào lớp mới, khả năng nhận diện helmet và no-helmet có thể suy giảm; do đó kết quả cần được đánh giá theo từng lớp trên cùng tập dữ liệu phù hợp.')
table('Bảng 1.2. So sánh phương án ban đầu và hướng phát triển', ['Tiêu chí', 'Hai lớp ban đầu', 'Ba lớp và liên kết'], [
    ('Lớp nhận diện', 'helmet, no-helmet', 'helmet, no-helmet, bike'),
    ('Thông tin xe máy', 'Không có vùng xe riêng', 'Có vị trí xe để liên kết'),
    ('Căn cứ ghi nhận', 'Chủ yếu dựa vào no-helmet', 'Kết hợp trạng thái đầu và quan hệ với xe'),
    ('Rủi ro chính', 'Báo nhầm người đi bộ', 'Liên kết nhầm hoặc bỏ sót khi không thấy xe'),
], [3.6, 5.5, 7.0])
para('Lớp xe máy chỉ cung cấp dữ liệu cho bước kiểm tra ngữ cảnh, không tự bảo đảm hệ thống chính xác hơn. Nếu xe bị che khuất hoặc mô hình bỏ sót xe, trường hợp thật có thể không được xác nhận. Vì vậy, so sánh hai phương án cần thực hiện trên cùng dữ liệu và cùng quy tắc đếm, xét đồng thời mức giảm báo nhầm người đi bộ, mức bỏ sót và tốc độ xử lý.')
h2('1.4.3. Liên kết vùng đầu với xe máy và theo dõi video')
para('Sau khi nhận diện, hệ thống kiểm tra vị trí tương đối, khoảng cách và mức phù hợp giữa vùng đầu với khu vực người ngồi trên xe. Một vùng no-helmet chỉ được xem là ứng viên ghi nhận khi có xe máy liên quan theo tiêu chí đã xác định. Sự xuất hiện đồng thời trong cùng ảnh không đủ: người đi bộ ở một phía của ảnh không thể được gắn với xe ở phía khác, còn người đứng sát xe vẫn có thể không phải người sử dụng xe. Khi có nhiều xe, hệ thống cần lựa chọn mối liên kết phù hợp, tránh gán một vùng đầu tùy ý cho phương tiện gần nhất.')
para('Đối với video, phương pháp theo dõi dựa trên tâm đối tượng được sử dụng để duy trì mã qua các khung hình. Cùng một trường hợp được ghi nhận một lần trong phạm vi tệp thay vì lặp lại theo số khung xuất hiện. Trạng thái no-helmet được xem xét qua nhiều lần phát hiện nhằm giảm tác động của dự đoán sai nhất thời. Ngưỡng xác nhận phải cân bằng: quá ít lần có thể làm tăng báo nhầm; quá nhiều lần có thể bỏ sót đối tượng chỉ xuất hiện ngắn. Phần phương pháp và tham số triển khai được trình bày ở Chương 2 và chương thiết kế hệ thống.')
h2('1.4.4. Tích hợp ứng dụng web và khai thác kết quả')
para('Ứng dụng web tổ chức dữ liệu theo dự án, địa điểm và phiên phân tích. Người dùng chọn ngày, khung giờ, tải ảnh hoặc video, theo dõi tiến độ và xem kết quả có đánh dấu. Hình ảnh minh chứng được lưu cùng tệp nguồn và thời điểm để người dùng kiểm tra một bản ghi trước khi sử dụng. Thống kê số lượng và tỷ lệ theo địa điểm, thời gian hỗ trợ nhận biết nơi, lúc cần quan tâm trong dữ liệu đã xử lý.')
para('Chất lượng thống kê phụ thuộc vào chất lượng nhận diện, liên kết và tracking. Khi so sánh địa điểm, cần xem xét thời lượng quan sát, lưu lượng xe, dữ liệu trùng lặp và điều kiện ghi hình. Hệ thống cung cấp thông tin tham khảo cho tuyên truyền và lựa chọn nơi ưu tiên giám sát; việc bố trí hoạt động thực tế vẫn cần con người xem xét bối cảnh và kiểm tra kết quả.')
h2('1.4.5. Các tình huống khó và giới hạn cần đánh giá')
para('Người đi bộ gần xe là tình huống kiểm thử quan trọng: khoảng cách nhỏ trong ảnh có thể khiến vùng đầu được liên kết nhầm với phương tiện. Xe chở nhiều người cũng đặt ra giới hạn về vai trò; mô hình hiện chưa phân biệt riêng người lái và người ngồi sau. Nếu người lái đội mũ nhưng người ngồi sau không đội mũ, quan hệ với xe không đủ để kết luận người điều khiển không đội mũ. Báo cáo cần tách rõ kết quả “vùng đầu không mũ liên quan đến xe” với kết luận về người lái khi chưa có căn cứ bổ sung.')
para('Các trường hợp vùng đầu nhỏ, ảnh nhòe, thiếu sáng và phương tiện bị che khuất có thể làm giảm khả năng phát hiện. Bỏ sót lớp bike khiến bước liên kết không xác nhận được một vùng no-helmet thực sự liên quan đến xe. Trong video, chuyển động nhanh, nhiều xe đi sát nhau hoặc che khuất kéo dài có thể làm mất hay đổi mã theo dõi, dẫn đến đếm trùng hoặc gộp nhầm đối tượng. Vì vậy, đánh giá cần tách phát hiện đúng, báo nhầm, bỏ sót và sai số đếm; số tổng bằng đối chứng chưa chứng minh mọi trường hợp đều được nhận diện đúng.')
para('Bộ kiểm thử cần bao gồm cảnh chỉ có người đi bộ, người đi bộ gần xe, xe chở nhiều người, xe bị che khuất và cảnh giao thông đông. So sánh mô hình hai lớp với phương án ba lớp phải sử dụng cùng dữ liệu và quy tắc đối chứng. Ngoài chỉ số nhận diện từng lớp, cần kiểm tra kết quả sau liên kết và tracking để xác định liệu phương án mới giảm báo nhầm mà vẫn duy trì khả năng phát hiện đối tượng cần tìm hay không.')

h1('1.5. Kết luận chương')
para('Chương 1 đã xác định nhu cầu hỗ trợ ghi nhận trường hợp không đội mũ bảo hiểm từ ảnh, video, lưu minh chứng và tổng hợp kết quả theo địa điểm, thời gian. Hạn chế của mô hình hai lớp là nhãn no-helmet không cho biết vùng đầu có thuộc người đi xe máy hay không, nên việc ghi nhận trực tiếp có thể nhầm người đi bộ. Hướng phát triển của đề tài là bổ sung lớp bike, kiểm tra quan hệ giữa vùng đầu và xe máy, kết hợp tracking và xác nhận trạng thái trong video.')
para('Các giới hạn về người đi bộ gần xe, phân biệt người lái với người ngồi sau, che khuất và sai số đếm cần được kiểm chứng bằng thực nghiệm. Kết quả thống kê chỉ có ý nghĩa trong phạm vi dữ liệu đã phân tích và phải được người dùng kiểm tra khi dùng cho tuyên truyền hoặc tham khảo bố trí giám sát. Những nội dung này là cơ sở để trình bày lý thuyết, công nghệ và thiết kế hệ thống ở các chương tiếp theo.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.save(OUT)
print(OUT)
